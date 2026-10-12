
import io

import docx
import fitz
import pytest

from agents.parser_agent import extract_text


class FakeFile(io.BytesIO):
    """File-like object with a filename, like an uploaded file."""

    def __init__(self, content=b"", name="sample.txt"):
        super().__init__(content)
        self.name = name


def test_none_file_returns_empty_string():
    assert extract_text(None) == ""


def test_txt_extracts_text():
    file = FakeFile(b"Hello RecruitIQ", "resume.txt")

    assert extract_text(file) == "Hello RecruitIQ"


def test_txt_invalid_utf8_is_ignored():
    file = FakeFile(b"Hello\xffWorld", "resume.txt")

    assert extract_text(file) == "HelloWorld"


def test_empty_txt_returns_empty_string():
    file = FakeFile(b"", "empty.txt")

    assert extract_text(file) == ""


def test_unsupported_extension_returns_empty_string():
    file = FakeFile(b"some content", "resume.csv")

    assert extract_text(file) == ""


def test_docx_extracts_multiple_paragraphs():
    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Python Developer")
    document.add_paragraph("Machine Learning Engineer")
    document.save(buffer)

    file = FakeFile(buffer.getvalue(), "resume.docx")

    result = extract_text(file)

    assert result == "Python Developer\nMachine Learning Engineer"


def test_empty_pdf_returns_empty_string():
    file = FakeFile(b"", "empty.pdf")

    assert extract_text(file) == ""


def test_pdf_extracts_text():
    buffer = io.BytesIO()

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "AI Engineer Resume")
    pdf_bytes = document.tobytes()
    document.close()

    file = FakeFile(pdf_bytes, "resume.pdf")

    assert extract_text(file) == "AI Engineer Resume"


def test_malformed_pdf_returns_empty_string(capsys):
    file = FakeFile(b"not a real PDF", "broken.pdf")

    assert extract_text(file) == ""
    assert "EXTRACTION ERROR:" in capsys.readouterr().out


def test_malformed_docx_returns_empty_string(capsys):
    file = FakeFile(b"not a real DOCX", "broken.docx")

    assert extract_text(file) == ""
    assert "EXTRACTION ERROR:" in capsys.readouterr().out


def test_file_is_rewound_before_reading():
    file = FakeFile(b"Resume text", "resume.txt")
    file.seek(5)

    assert extract_text(file) == "Resume text"
